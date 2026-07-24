from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.audit_log import write_audit_log
from src.database import DEFAULT_DB_PATH, PROJECT_ROOT, get_connection


DEFAULT_EXPORT_DIR = PROJECT_ROOT / "data" / "exports"
DETAIL_COLUMNS = ["物料编码", "物料名称", "风险等级", "推荐结论", "风险原因", "建议动作"]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial")
BODY_FONT = Font(name="Arial")
ZEBRA_FILL_1 = PatternFill(fill_type="solid", fgColor="FFFFFF")
ZEBRA_FILL_2 = PatternFill(fill_type="solid", fgColor="F7F9FC")
HIGH_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
MEDIUM_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
LOW_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9DEE7"),
    right=Side(style="thin", color="D9DEE7"),
    top=Side(style="thin", color="D9DEE7"),
    bottom=Side(style="thin", color="D9DEE7"),
)


def load_report_results(batch_id: int, db_path: Path | str = DEFAULT_DB_PATH) -> pd.DataFrame:
    with get_connection(db_path) as conn:
        rows = conn.execute(
            """
            SELECT
                material_code AS 物料编码,
                material_name AS 物料名称,
                risk_level AS 风险等级,
                recommendation AS 推荐结论,
                risk_reasons AS 风险原因,
                suggested_action AS 建议动作
            FROM risk_results
            WHERE batch_id = ?
            ORDER BY id
            """,
            (batch_id,),
        ).fetchall()
    return pd.DataFrame([dict(row) for row in rows], columns=DETAIL_COLUMNS)


def build_report_summary(batch_id: int, detail_df: pd.DataFrame) -> list[tuple[str, object]]:
    risk_counts = detail_df["风险等级"].value_counts().to_dict() if not detail_df.empty else {}
    return [
        ("BOM批次ID", batch_id),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("物料总数", int(len(detail_df))),
        ("高风险数量", int(risk_counts.get("高", 0))),
        ("中风险数量", int(risk_counts.get("中", 0))),
        ("低风险数量", int(risk_counts.get("低", 0))),
    ]


def _apply_header_style(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _risk_fill(risk_level: str):
    if risk_level == "高":
        return HIGH_FILL
    if risk_level == "中":
        return MEDIUM_FILL
    if risk_level == "低":
        return LOW_FILL
    return None


def _apply_table_style(ws, long_text_columns: set[int] | None = None) -> None:
    long_text_columns = long_text_columns or set()
    _apply_header_style(ws)
    for row_idx in range(2, ws.max_row + 1):
        fill = ZEBRA_FILL_1 if row_idx % 2 == 0 else ZEBRA_FILL_2
        risk_level = ws.cell(row=row_idx, column=3).value if ws.max_column >= 3 else None
        risk_level_fill = _risk_fill(risk_level)
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = BODY_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col_idx in long_text_columns)
            if col_idx == 3 and risk_level_fill:
                cell.fill = risk_level_fill


def _auto_fit_columns(ws, wrap_columns: set[int] | None = None) -> None:
    wrap_columns = wrap_columns or set()
    for col_idx in range(1, ws.max_column + 1):
        max_len = max(
            (len(str(ws.cell(row=row, column=col_idx).value or "")) for row in range(1, ws.max_row + 1)),
            default=0,
        )
        width = min(max(max_len + 2, 8), 60)
        if col_idx in wrap_columns:
            width = min(max(width, 40), 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _append_dataframe(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append([row[column] for column in df.columns])


def _build_workbook(batch_id: int, detail_df: pd.DataFrame) -> Workbook:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "风险摘要"
    detail_ws = wb.create_sheet("评分明细")
    not_recommended_ws = wb.create_sheet("不推荐使用物料")

    summary_ws.append(["指标", "值"])
    for item in build_report_summary(batch_id, detail_df):
        summary_ws.append(list(item))
    _apply_table_style(summary_ws)
    for row_idx in range(2, summary_ws.max_row + 1):
        metric = summary_ws.cell(row=row_idx, column=1).value
        value_cell = summary_ws.cell(row=row_idx, column=2)
        if metric == "高风险数量":
            value_cell.fill = HIGH_FILL
        elif metric == "中风险数量":
            value_cell.fill = MEDIUM_FILL
        elif metric == "低风险数量":
            value_cell.fill = LOW_FILL
        if "数量" in str(metric) or metric == "物料总数":
            value_cell.number_format = "#,##0"
    _auto_fit_columns(summary_ws)

    _append_dataframe(detail_ws, detail_df)
    _apply_table_style(detail_ws, long_text_columns={5, 6})
    _auto_fit_columns(detail_ws, wrap_columns={5, 6})

    not_recommended_df = detail_df[detail_df["推荐结论"] == "不推荐使用"]
    _append_dataframe(not_recommended_ws, not_recommended_df)
    _apply_table_style(not_recommended_ws, long_text_columns={5, 6})
    _auto_fit_columns(not_recommended_ws, wrap_columns={5, 6})

    return wb


def export_review_report(
    batch_id: int,
    export_dir: Path | str = DEFAULT_EXPORT_DIR,
    db_path: Path | str = DEFAULT_DB_PATH,
    operator: str | None = None,
) -> Path:
    detail_df = load_report_results(batch_id, db_path)
    export_path = Path(export_dir)
    export_path.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = export_path / f"BOM评审报告_批次{batch_id}_{timestamp}.xlsx"

    wb = _build_workbook(batch_id, detail_df)
    wb.save(report_path)

    write_audit_log(
        operation_type="导出报告",
        operation_object=report_path.name,
        operation_result="成功",
        operator=operator,
        message=f"导出批次{batch_id}评审报告",
        db_path=db_path,
    )
    return report_path
