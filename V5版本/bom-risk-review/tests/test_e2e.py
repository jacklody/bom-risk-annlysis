from pathlib import Path

from openpyxl import load_workbook

from src.database import initialize_database
from src.importers import read_bom_excel, read_supply_excel, save_bom_import, save_supply_import
from src.report_exporter import export_review_report
from src.scoring_service import run_scoring_for_batch


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE_BOM = PROJECT_ROOT / "data" / "samples" / "sample_bom.xlsx"
SAMPLE_SUPPLY = PROJECT_ROOT / "data" / "samples" / "sample_supply.xlsx"


def test_e2e_full_workflow_from_empty_database(tmp_path):
    db_path = tmp_path / "app.db"
    export_dir = tmp_path / "exports"

    # 1. 从空数据库开始，初始化数据库
    initialize_database(db_path)

    # 2. 导入样例 BOM
    bom_df = read_bom_excel(SAMPLE_BOM, db_path=db_path)
    batch_id = save_bom_import(bom_df, file_name="sample_bom.xlsx", db_path=db_path, operator="测试用户")

    # 3. 导入样例供应数据
    supply_df = read_supply_excel(SAMPLE_SUPPLY, db_path=db_path)
    save_supply_import(supply_df, file_name="sample_supply.xlsx", db_path=db_path, operator="测试用户")

    # 4. 执行风险评分
    results, summary = run_scoring_for_batch(batch_id, db_path=db_path, operator="测试用户")

    # 5. 验证评分产生高/中/低风险结果
    risk_levels = set(results["风险等级"].unique())
    assert "高" in risk_levels
    assert "中" in risk_levels
    assert "低" in risk_levels
    assert summary["物料总数"] == len(results)
    assert summary["高风险数量"] + summary["中风险数量"] + summary["低风险数量"] == summary["物料总数"]

    # 6. 导出评审报告
    report_path = export_review_report(batch_id, export_dir=export_dir, db_path=db_path, operator="测试用户")

    # 7. 验证报告包含三张 Sheet
    wb = load_workbook(report_path)
    assert wb.sheetnames == ["风险摘要", "评分明细", "不推荐使用物料"]

    # 8. 验证风险摘要统计与评分结果一致
    ws = wb["风险摘要"]
    summary_values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert summary_values["物料总数"] == summary["物料总数"]
    assert summary_values["高风险数量"] == summary["高风险数量"]
    assert summary_values["中风险数量"] == summary["中风险数量"]
    assert summary_values["低风险数量"] == summary["低风险数量"]

    # 9. 评分明细记录数等于 BOM 有效物料数量
    detail_ws = wb["评分明细"]
    assert detail_ws.max_row - 1 == len(results)

    # 10. 不推荐使用物料只包含高风险
    not_recommended_ws = wb["不推荐使用物料"]
    if not_recommended_ws.max_row > 1:
        for r in range(2, not_recommended_ws.max_row + 1):
            assert not_recommended_ws.cell(row=r, column=4).value == "不推荐使用"
