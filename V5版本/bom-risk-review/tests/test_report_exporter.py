import sqlite3

from openpyxl import load_workbook

from src.database import initialize_database
from src.importers import read_bom_excel, read_supply_excel, save_bom_import, save_supply_import
from src.report_exporter import export_review_report
from src.sample_data import create_sample_files
from src.scoring_service import run_scoring_for_batch


def prepare_scored_batch(tmp_path):
    db_path = tmp_path / "app.db"
    initialize_database(db_path)
    sample_dir = tmp_path / "samples"
    export_dir = tmp_path / "exports"
    bom_path, supply_path = create_sample_files(sample_dir)
    bom_df = read_bom_excel(bom_path)
    supply_df = read_supply_excel(supply_path)
    batch_id = save_bom_import(bom_df, file_name="sample_bom.xlsx", db_path=db_path, operator="测试用户")
    save_supply_import(supply_df, file_name="sample_supply.xlsx", db_path=db_path, operator="测试用户")
    results, summary = run_scoring_for_batch(batch_id, db_path=db_path, operator="测试用户")
    return db_path, export_dir, batch_id, len(results), summary


def test_export_review_report_creates_required_workbook_sheets(tmp_path):
    db_path, export_dir, batch_id, _, _ = prepare_scored_batch(tmp_path)

    report_path = export_review_report(batch_id, export_dir=export_dir, db_path=db_path, operator="测试用户")

    assert report_path.exists()
    assert report_path.suffix == ".xlsx"
    wb = load_workbook(report_path)
    assert wb.sheetnames == ["风险摘要", "评分明细", "不推荐使用物料"]


def test_export_review_report_summary_matches_scoring_results(tmp_path):
    db_path, export_dir, batch_id, row_count, summary = prepare_scored_batch(tmp_path)

    report_path = export_review_report(batch_id, export_dir=export_dir, db_path=db_path, operator="测试用户")

    wb = load_workbook(report_path, data_only=True)
    ws = wb["风险摘要"]
    summary_values = {ws.cell(row=row, column=1).value: ws.cell(row=row, column=2).value for row in range(2, ws.max_row + 1)}

    assert summary_values["物料总数"] == row_count
    assert summary_values["高风险数量"] == summary["高风险数量"]
    assert summary_values["中风险数量"] == summary["中风险数量"]
    assert summary_values["低风险数量"] == summary["低风险数量"]


def test_export_review_report_detail_and_not_recommended_sheets(tmp_path):
    db_path, export_dir, batch_id, row_count, _ = prepare_scored_batch(tmp_path)

    report_path = export_review_report(batch_id, export_dir=export_dir, db_path=db_path, operator="测试用户")

    wb = load_workbook(report_path, data_only=True)
    detail_ws = wb["评分明细"]
    not_recommended_ws = wb["不推荐使用物料"]

    assert detail_ws.max_row - 1 == row_count
    recommendations = [not_recommended_ws.cell(row=row, column=4).value for row in range(2, not_recommended_ws.max_row + 1)]
    assert recommendations
    assert set(recommendations) == {"不推荐使用"}


def test_export_review_report_writes_audit_log(tmp_path):
    db_path, export_dir, batch_id, _, _ = prepare_scored_batch(tmp_path)

    report_path = export_review_report(batch_id, export_dir=export_dir, db_path=db_path, operator="测试用户")

    with sqlite3.connect(db_path) as conn:
        log = conn.execute("SELECT operation_type, operation_object, operation_result, operator FROM audit_logs ORDER BY id DESC LIMIT 1").fetchone()

    assert log == ("导出报告", report_path.name, "成功", "测试用户")
