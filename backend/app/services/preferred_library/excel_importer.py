import json
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from app.models.preferred import PreferredPart
from app.services.preferred_library.normalizer import normalize_brand, normalize_mpn, normalize_text, parse_preferred_level


class PreferredLibraryImporter:
    def __init__(self, db: Session):
        self.db = db

    def import_file(self, file_path: str, source_file: str | None = None) -> dict[str, Any]:
        wb = load_workbook(file_path, data_only=True)
        source_name = source_file or Path(file_path).name
        self.db.query(PreferredPart).filter(PreferredPart.source_file == source_name).delete()
        imported = 0
        strong = 0
        medium = 0
        pending = 0
        sheets = 0
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            headers = [normalize_text(v) for v in rows[0]]
            if not any(headers):
                continue
            sheets += 1
            for idx, row in enumerate(rows[1:], start=2):
                data = self._row_to_dict(headers, row)
                if not self._has_part_identity(data):
                    continue
                part = self._build_part(data, source_name, ws.title, idx)
                self.db.add(part)
                imported += 1
                if part.preferred_level == "strong_recommended":
                    strong += 1
                if part.preferred_level == "medium_recommended":
                    medium += 1
                if part.control_status == "pending_review":
                    pending += 1
        self.db.commit()
        return {
            "source_file": source_name,
            "sheets": sheets,
            "imported_parts": imported,
            "strong_recommended_count": strong,
            "medium_recommended_count": medium,
            "pending_review_count": pending,
        }

    def _row_to_dict(self, headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
        return {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}

    def _get(self, data: dict[str, Any], names: list[str]) -> str:
        for name in names:
            if name in data and normalize_text(data[name]):
                return normalize_text(data[name])
        return ""

    def _has_part_identity(self, data: dict[str, Any]) -> bool:
        return bool(self._get(data, ["物料号", "型号", "厂家型号", "名称", "品牌"]))

    def _build_part(self, data: dict[str, Any], source_file: str, sheet: str, row: int) -> PreferredPart:
        brand = self._get(data, ["品牌", "厂家", "制造商"])
        mpn = self._get(data, ["型号", "规格型号", "厂家型号"])
        status = self._get(data, ["通用化状态", "状态", "推荐状态", "依据实际更改"])
        level, control = parse_preferred_level(status)
        spec_keys = [k for k in data.keys() if k not in {"品牌", "名称", "通用化状态", "物料号", "型号", "厂家型号", "封装形式", "应用方案", "成本", "未税", "含税"}]
        key_params = {k: normalize_text(data[k]) for k in spec_keys if normalize_text(data[k])}
        cost_text = self._get(data, ["成本", "未税", "含税"])
        return PreferredPart(
            source_file=source_file,
            source_sheet=sheet,
            source_row=row,
            category_level_2=sheet,
            category_level_3=sheet,
            brand=brand,
            normalized_brand=normalize_brand(brand),
            name=self._get(data, ["名称", "品名", "物料名称"]),
            part_number=self._get(data, ["物料号", "内部物料号"]),
            mpn=mpn,
            normalized_mpn=normalize_mpn(mpn),
            manufacturer_mpn=self._get(data, ["厂家型号"]),
            package=self._get(data, ["封装形式", "封装"]),
            mounting_type=self._get(data, ["加工方式"]),
            size=self._get(data, ["尺寸"]),
            spec_text=" ".join(v for v in [mpn, self._get(data, ["封装形式", "尺寸"]), self._get(data, ["电压", "容量", "电流", "电感量"])] if v),
            key_params_json=json.dumps(key_params, ensure_ascii=False),
            application_scenarios=self._get(data, ["应用方案", "应用", "方案"]),
            preferred_level=level,
            control_status=control,
            cost=float(cost_text) if cost_text and cost_text.replace(".", "", 1).isdigit() else None,
            cost_unit="CNY" if cost_text else None,
            recommendation_reason=status,
            raw_row_json=json.dumps({k: normalize_text(v) for k, v in data.items()}, ensure_ascii=False),
        )
