from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _text(value) -> str:
    if value is None:
        return ""
    return escape(str(value))


def _risk_label(level: str | None) -> str:
    return {"red": "高风险", "yellow": "关注", "green": "安全", "gray": "无数据"}.get(level or "gray", "无数据")


def _render_rows(items: list[dict]) -> str:
    if not items:
        return "<tr><td colspan='8'>暂无高风险物料</td></tr>"
    rows = []
    for item in items:
        evidence = "；".join(_text(event.get("title")) for event in item.get("market_evidence") or [])
        rows.append(
            "<tr>"
            f"<td>{_text(item.get('part_number'))}</td>"
            f"<td>{_text(item.get('brand'))}</td>"
            f"<td>{_text(item.get('category'))}</td>"
            f"<td>{_text(_risk_label(item.get('risk_level')))}</td>"
            f"<td>{_text(item.get('total_stock'))}</td>"
            f"<td>{_text(item.get('lead_time_days'))}</td>"
            f"<td>{_text(item.get('supply_sources_summary'))}</td>"
            f"<td>{evidence}</td>"
            "</tr>"
        )
    return "".join(rows)


def _render_actions(actions: dict) -> str:
    rows = []
    for action in actions.get("urgent") or []:
        rows.append(f"<li><strong>紧急：</strong>{_text(action)}</li>")
    for action in actions.get("watch") or []:
        rows.append(f"<li><strong>关注：</strong>{_text(action)}</li>")
    if not rows:
        return "<li>暂无采购建议</li>"
    return "".join(rows)


def render_risk_report_html(report: dict) -> str:
    project = report.get("project") or {}
    summary = report.get("risk_summary") or {}
    items = report.get("high_risk_items") or []
    actions = report.get("actions") or {}
    title = _text(project.get("project_name") or project.get("filename") or "BOM风险报告")
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>{title}</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Arial, sans-serif; margin: 32px; color: #1f2937; background: #f8fafc; }}
.report {{ max-width: 1180px; margin: 0 auto; background: #fff; padding: 32px; border-radius: 16px; box-shadow: 0 10px 30px rgba(15, 23, 42, 0.08); }}
h1 {{ margin: 0 0 8px; font-size: 28px; }}
h2 {{ margin-top: 28px; font-size: 20px; border-left: 4px solid #2563eb; padding-left: 10px; }}
.meta, .cards {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin-top: 16px; }}
.card {{ background: #f1f5f9; border-radius: 12px; padding: 14px; }}
.card span {{ display: block; color: #64748b; font-size: 13px; }}
.card strong {{ display: block; margin-top: 6px; font-size: 22px; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 12px; font-size: 13px; }}
th, td {{ border: 1px solid #e5e7eb; padding: 10px; text-align: left; vertical-align: top; }}
th {{ background: #f8fafc; }}
.notice {{ margin-top: 20px; color: #64748b; font-size: 12px; }}
@media print {{ body {{ margin: 0; background: #fff; }} .report {{ box-shadow: none; border-radius: 0; }} }}
</style>
</head>
<body>
<main class="report">
<h1>{title}</h1>
<p>项目风险报告 · 生成时间：{_text(report.get('last_analysis_at'))}</p>
<section class="meta">
<div class="card"><span>文件名</span><strong>{_text(project.get('filename'))}</strong></div>
<div class="card"><span>物料总数</span><strong>{_text(summary.get('total_items'))}</strong></div>
<div class="card"><span>风险率</span><strong>{_text(summary.get('risk_rate'))}%</strong></div>
<div class="card"><span>紧急风险</span><strong>{'是' if summary.get('has_urgent') else '否'}</strong></div>
</section>
<section class="cards">
<div class="card"><span>高风险</span><strong>{_text(summary.get('red'))}</strong></div>
<div class="card"><span>关注</span><strong>{_text(summary.get('yellow'))}</strong></div>
<div class="card"><span>安全</span><strong>{_text(summary.get('green'))}</strong></div>
<div class="card"><span>无数据</span><strong>{_text(summary.get('gray'))}</strong></div>
</section>
<h2>高风险物料与证据</h2>
<table>
<thead><tr><th>型号</th><th>品牌</th><th>品类</th><th>风险</th><th>库存</th><th>交期天数</th><th>供应证据</th><th>行情证据</th></tr></thead>
<tbody>{_render_rows(items)}</tbody>
</table>
<h2>采购建议</h2>
<ul>{_render_actions(actions)}</ul>
<p class="notice">本报告由系统已有 BOM、供应与行情证据生成，不包含密钥、环境变量或内部调试信息。可使用浏览器打印功能另存为 PDF。</p>
</main>
</body>
</html>"""


def _plain(value):
    return "" if value is None else value


def _pct(value):
    if value is None or value == "":
        return ""
    try:
        return float(value)
    except (TypeError, ValueError):
        return ""


def _append_table(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for row in rows:
        ws.append(row)


def _style_sheet(ws, wrap_columns: set[int] | None = None, money_columns: set[int] | None = None, pct_columns: set[int] | None = None):
    wrap_columns = wrap_columns or set()
    money_columns = money_columns or set()
    pct_columns = pct_columns or set()
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    zebra_fill = PatternFill(fill_type="solid", fgColor="F7F9FC")
    border = Border(
        left=Side(style="thin", color="D9DEE7"),
        right=Side(style="thin", color="D9DEE7"),
        top=Side(style="thin", color="D9DEE7"),
        bottom=Side(style="thin", color="D9DEE7"),
    )
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            if cell.row > 1 and cell.row % 2 == 0:
                cell.fill = zebra_fill
            if cell.column in wrap_columns:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
            if cell.column in money_columns and cell.row > 1:
                cell.number_format = "¥#,##0.00"
            elif cell.column in pct_columns and cell.row > 1:
                cell.number_format = "0.0%"
            elif isinstance(cell.value, int) and cell.row > 1:
                cell.number_format = "#,##0"
    for col_idx in range(1, ws.max_column + 1):
        max_len = max((len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)), default=0)
        width = min(max(max_len + 2, 10), 60)
        if col_idx in wrap_columns:
            width = max(width, 32)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"


def _summary_rows(report: dict) -> list[list]:
    project = report.get("project") or {}
    summary = report.get("risk_summary") or {}
    return [
        ["项目名称", _plain(project.get("project_name") or project.get("filename"))],
        ["文件名", _plain(project.get("filename"))],
        ["物料总数", summary.get("total_items") or 0],
        ["红色物料", summary.get("red") or 0],
        ["黄色物料", summary.get("yellow") or 0],
        ["绿色物料", summary.get("green") or 0],
        ["灰色物料", summary.get("gray") or 0],
        ["风险率(%)", summary.get("risk_rate") or 0],
        ["生成时间", _plain(report.get("last_analysis_at"))],
    ]


def _item_rows(report: dict) -> list[list]:
    rows = []
    for item in report.get("high_risk_items") or []:
        rows.append([
            _plain(item.get("part_number")),
            _plain(item.get("brand")),
            _plain(item.get("category")),
            _risk_label(item.get("risk_level")),
            item.get("risk_score") or 0,
            item.get("total_stock") or 0,
            item.get("supplier_count") or 0,
            item.get("lead_time_days") or "",
            item.get("unit_price") or "",
            _plain(item.get("supply_sources_summary")),
            _plain(item.get("procurement_suggestion") or item.get("suggestion")),
        ])
    return rows


def _evidence_rows(report: dict) -> list[list]:
    rows = []
    for item in report.get("high_risk_items") or []:
        part_number = _plain(item.get("part_number"))
        for event in item.get("market_evidence") or []:
            rows.append([
                part_number,
                _plain(event.get("title")),
                _plain(event.get("source_name")),
                _plain(event.get("event_type")),
                _plain(event.get("severity")),
                _pct(event.get("confidence")),
                event.get("price_impact_pct") or "",
                event.get("lead_time_impact_weeks") or "",
            ])
    return rows


def render_risk_report_xlsx(report: dict) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "项目摘要"
    for row in _summary_rows(report):
        summary_ws.append(row)
    _style_sheet(summary_ws)
    summary_ws["B4"].fill = PatternFill(fill_type="solid", fgColor="FDE2E2")
    summary_ws["B5"].fill = PatternFill(fill_type="solid", fgColor="FFF4D6")

    items_ws = wb.create_sheet("高风险物料")
    _append_table(
        items_ws,
        ["物料型号", "品牌", "品类", "风险等级", "风险分数", "库存数量", "供应商数", "交期天数", "单价(元)", "供应摘要", "采购建议"],
        _item_rows(report),
    )
    _style_sheet(items_ws, wrap_columns={10, 11}, money_columns={9})

    evidence_ws = wb.create_sheet("行情证据")
    _append_table(
        evidence_ws,
        ["物料型号", "事件标题", "来源", "事件类型", "严重度", "置信度", "价格影响(%)", "交期影响(周)"],
        _evidence_rows(report),
    )
    _style_sheet(evidence_ws, wrap_columns={2}, pct_columns={6})

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
