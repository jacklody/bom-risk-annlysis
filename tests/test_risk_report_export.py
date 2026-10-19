from datetime import datetime, timezone
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import SessionLocal, init_db
from app.main import app
from app.models.bom import BomItem, BomProject, RiskLevel
from app.models.event import EventType, MarketEvent, Severity

client = TestClient(app)


def _clean_report_data():
    init_db()
    db = SessionLocal()
    db.query(BomItem).delete()
    db.query(BomProject).delete()
    db.query(MarketEvent).delete()
    db.commit()
    db.close()


def _seed_report_project():
    db = SessionLocal()
    try:
        project = BomProject(
            filename="report-demo.xlsx",
            project_name="报告导出演示项目",
            status="active",
            total_items=1,
            items_analyzed=1,
            items_red=1,
            items_yellow=0,
            items_green=0,
            items_gray=0,
            last_analysis_at=datetime.now(timezone.utc),
        )
        db.add(project)
        db.commit()
        db.refresh(project)
        event = MarketEvent(
            source_name="供应商通知",
            source_url="test://report-event",
            raw_title="STM32 MCU 交期拉长",
            raw_text="STM32 MCU 交期拉长，建议提前锁单",
            brand="ST",
            category="MCU",
            event_type=EventType.LEAD_TIME,
            severity=Severity.HIGH,
            confidence=0.9,
            is_active=1,
        )
        db.add(event)
        db.commit()
        db.refresh(event)
        item = BomItem(
            bom_upload_id=project.id,
            row_number=1,
            raw_part_number="STM32F103C8T6",
            raw_brand="ST",
            raw_category="MCU",
            normalized_category="MCU",
            risk_level=RiskLevel.RED,
            risk_score=95,
            risk_reason="交期拉长并存在缺货风险",
            matched_event_ids=str(event.id),
            total_stock=0,
            supplier_count=1,
            lead_time_days=56,
            unit_price=12.5,
            lifecycle_status="Active",
            supply_sources_summary="lcsc:库存0; mouser:交期56天",
            supply_risk_tags="low_stock,long_lead_time",
            procurement_suggestion="建议锁单并准备替代料。",
        )
        db.add(item)
        db.commit()
        return project.id
    finally:
        db.close()


def test_risk_report_export_html_contains_project_supply_market_and_actions():
    _clean_report_data()
    project_id = _seed_report_project()

    resp = client.get(f"/api/risk/report/{project_id}/export.html")

    assert resp.status_code == 200
    assert "text/html" in resp.headers["content-type"]
    html = resp.text
    assert "报告导出演示项目" in html
    assert "STM32F103C8T6" in html
    assert "lcsc:库存0" in html
    assert "STM32 MCU 交期拉长" in html
    assert "建议锁单" in html
    assert "window.__" not in html


def test_risk_report_export_html_returns_404_for_missing_project():
    _clean_report_data()

    resp = client.get("/api/risk/report/999999/export.html")

    assert resp.status_code == 404


def test_risk_report_export_xlsx_contains_summary_items_and_market_evidence():
    _clean_report_data()
    project_id = _seed_report_project()

    resp = client.get(f"/api/risk/report/{project_id}/export.xlsx")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]

    wb = load_workbook(BytesIO(resp.content), data_only=False)
    assert wb.sheetnames == ["项目摘要", "高风险物料", "行情证据"]

    summary = wb["项目摘要"]
    assert summary["A1"].value == "项目名称"
    assert summary["B1"].value == "报告导出演示项目"
    assert summary["A4"].value == "红色物料"
    assert summary["B4"].value == 1

    items = wb["高风险物料"]
    assert items["A1"].value == "物料型号"
    assert items["A2"].value == "STM32F103C8T6"
    assert items["I1"].value == "单价(元)"
    assert items["I2"].value == 12.5
    assert items["K2"].value == "建议锁单并准备替代料。"

    evidence = wb["行情证据"]
    assert evidence["A1"].value == "物料型号"
    assert evidence["B2"].value == "STM32 MCU 交期拉长"
    assert evidence["C2"].value == "供应商通知"
    assert evidence["F2"].value == 0.9


def test_risk_report_export_xlsx_returns_404_for_missing_project():
    _clean_report_data()

    resp = client.get("/api/risk/report/999999/export.xlsx")

    assert resp.status_code == 404


# ── CSV 导出 ──────────────────────────────────────────────

def test_risk_report_export_csv_contains_all_items():
    _clean_report_data()
    project_id = _seed_report_project()

    resp = client.get(f"/api/risk/report/{project_id}/export.csv")

    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]

    csv_text = resp.text
    assert "行号" in csv_text
    assert "物料编号" in csv_text
    assert "风险等级" in csv_text
    assert "v3.2风险等级" in csv_text
    assert "关联事件标题" in csv_text
    assert "STM32F103C8T6" in csv_text
    assert "STM32 MCU 交期拉长" in csv_text
    assert "供应商通知" in csv_text


def test_risk_report_export_csv_returns_404_for_missing_project():
    _clean_report_data()

    resp = client.get("/api/risk/report/999999/export.csv")

    assert resp.status_code == 404


# ── 橙色风险 & 变更追踪 ──────────────────────────────────

def _seed_orange_risk_project():
    """创建含橙色风险和变更追踪数据的项目"""
    db = SessionLocal()
    try:
        project = BomProject(
            filename="orange-demo.xlsx",
            project_name="橙色风险测试项目",
            status="active",
            total_items=2,
            items_analyzed=2,
            items_red=0,
            items_orange=1,
            items_yellow=1,
            items_green=0,
            items_gray=0,
            new_orange_count=1,
            last_analysis_at=datetime.now(timezone.utc),
        )
        db.add(project)
        db.commit()
        db.refresh(project)

        event = MarketEvent(
            source_name="行业报告",
            source_url="test://orange-event",
            raw_title="MOSFET缺货预警",
            raw_text="行业MOSFET缺货预警，建议提前备货",
            brand="Infineon",
            category="MOSFET",
            event_type=EventType.SHORTAGE,
            severity=Severity.MEDIUM,
            confidence=0.7,
            is_active=1,
        )
        db.add(event)
        db.commit()
        db.refresh(event)

        changed_at = datetime(2026, 6, 28, 10, 0, 0, tzinfo=timezone.utc)

        # 橙色风险物料（带变更追踪）
        orange_item = BomItem(
            bom_upload_id=project.id,
            row_number=1,
            raw_part_number="IRLZ44N",
            raw_brand="Infineon",
            raw_category="MOSFET",
            normalized_category="MOSFET",
            risk_level=RiskLevel.ORANGE,
            risk_score=60.0,
            risk_reason="品类波及命中，中等关注",
            matched_event_ids=str(event.id),
            last_risk_level="yellow",
            risk_changed_at=changed_at,
            risk_trigger_event_ids=str(event.id),
            total_stock=500,
            supplier_count=2,
            lead_time_days=28,
            unit_price=3.2,
            lifecycle_status="Active",
            supply_sources_summary="lcsc:库存500; mouser:库存200",
            supply_risk_tags="medium_stock",
            procurement_suggestion="持续关注交期变化。",
        )
        db.add(orange_item)
        # 黄色风险物料（无变更）
        yellow_item = BomItem(
            bom_upload_id=project.id,
            row_number=2,
            raw_part_number="IRLZ34N",
            raw_brand="Infineon",
            raw_category="MOSFET",
            normalized_category="MOSFET",
            risk_level=RiskLevel.YELLOW,
            risk_score=25.0,
            risk_reason="普遍品类关注",
            matched_event_ids="",
            last_risk_level="yellow",
            risk_changed_at=None,
            risk_trigger_event_ids="",
            total_stock=1200,
            supplier_count=3,
            lead_time_days=14,
            unit_price=2.1,
            lifecycle_status="Active",
            supply_sources_summary="lcsc:库存1200",
            supply_risk_tags="",
            procurement_suggestion="",
        )
        db.add(yellow_item)
        db.commit()
        return project.id
    finally:
        db.close()


def test_report_contains_orange_risk_items():
    _clean_report_data()
    project_id = _seed_orange_risk_project()

    resp = client.get(f"/api/risk/report/{project_id}")

    assert resp.status_code == 200
    data = resp.json()
    summary = data["risk_summary"]
    assert summary["red"] == 0
    # 橙色出现在高风险物料中
    high_risk = data["high_risk_items"]
    assert len(high_risk) > 0
    orange_items = [it for it in high_risk if it["risk_level"] == "orange"]
    assert len(orange_items) >= 1
    assert orange_items[0]["part_number"] == "IRLZ44N"

    # CSV 导出中也应包含橙色
    csv_resp = client.get(f"/api/risk/report/{project_id}/export.csv")
    assert csv_resp.status_code == 200
    assert "IRLZ44N" in csv_resp.text
    assert "IRLZ34N" in csv_resp.text


def test_risk_change_tracking_fields_present():
    """验证变更追踪字段 last_risk_level/risk_changed_at 在 item 上持久化"""
    _clean_report_data()
    project_id = _seed_orange_risk_project()

    db = SessionLocal()
    try:
        orange = (
            db.query(BomItem)
            .filter(BomItem.bom_upload_id == project_id, BomItem.raw_part_number == "IRLZ44N")
            .first()
        )
        assert orange is not None
        assert orange.risk_level == RiskLevel.ORANGE
        assert orange.last_risk_level == "yellow"
        assert orange.risk_changed_at is not None
        assert orange.risk_trigger_event_ids is not None
    finally:
        db.close()
